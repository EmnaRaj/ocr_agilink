"""Export a fiche to Excel / PDF / CSV for download.

All three read the same normalised view of the fiche (its stored extraction +
work-order/product metadata), so the formats stay in sync.
"""

import csv
import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..models import Fiche

_BLUE = "0E5B75"  # Agilink brand petrol-blue
_LOGO = Path(__file__).resolve().parent.parent / "assets" / "logo.png"
_FARNESS = Path(__file__).resolve().parent.parent / "assets" / "farness.png"  # "powered by" footer

OP_HEADERS = ["#", "Opération", "Appl.", "Date", "Qté", "Début", "Fin", "Outillage", "Matricule"]
CTRL_HEADERS = ["Contrôle", "Méthode", "Résultat", "Matricule"]


def _data(fiche: Fiche) -> dict:
    ex = fiche.raw_extraction or {}
    h = ex.get("header") or {}
    hv = lambda k: (h.get(k) or {}).get("value")  # noqa: E731
    meta = ex.get("meta") or {}
    return {
        "ref_produit": fiche.work_order.product.ref_produit,
        "designation": fiche.work_order.product.designation,
        "n_of": fiche.work_order.n_of,
        "qte": hv("qte") if hv("qte") is not None else fiche.work_order.quantite,
        "annotation": hv("annotation_serie"),
        "statut": fiche.statut.value,
        "date_creation": fiche.date_creation,
        "confidence": meta.get("overall_confidence"),
        "operations": ex.get("operations") or [],
        "controls": ex.get("controls") or [],
        "validation": ex.get("validation") or [],
    }


def _t(field: object) -> str:
    v = (field or {}).get("value") if isinstance(field, dict) else None
    return str(v)[:5] if v else ""


def _op_cells(o: dict) -> list:
    fv = lambda k: (o.get(k) or {}).get("value")  # noqa: E731
    date = (o.get("date_op") or {}).get("raw_text") or fv("date_op") or ""
    appl = {True: "Oui", False: "Non"}.get(fv("applicable"), "")
    return [
        o.get("ordre"),
        o.get("nom_operation"),
        appl,
        date,
        fv("qte_realisee") if fv("qte_realisee") is not None else "",
        _t(o.get("heure_debut")),
        _t(o.get("heure_fin")),
        fv("outillage") or "",
        fv("matricule_operateur") or "",
    ]


def _ctrl_cells(c: dict) -> list:
    fv = lambda k: (c.get(k) or {}).get("value")  # noqa: E731
    meth = fv("methode")
    if isinstance(meth, dict):
        meth = meth.get("value")
    res = {True: "Conforme", False: "Non conforme"}.get(fv("resultat"), "")
    return [c.get("nom_operation"), meth or "", res, fv("matricule_operateur") or ""]


def file_stem(fiche: Fiche) -> str:
    d = _data(fiche)
    return f"fiche_{d['ref_produit']}_OF{d['n_of']}".replace(" ", "")


# --- Excel ------------------------------------------------------------------


def to_xlsx(fiche: Fiche) -> bytes:
    d = _data(fiche)
    wb = Workbook()
    ws = wb.active
    ws.title = "Fiche Suiveuse"
    white_bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=_BLUE)

    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value = "AGILINK GROUP — Fiche Suiveuse / Traçabilité Interne"
    c.font = Font(bold=True, size=14, color="FFFFFF")
    c.fill = fill
    c.alignment = Alignment(horizontal="center")

    row = 3
    for label, val in [
        ("Réf. Produit", d["ref_produit"]),
        ("Désignation", d["designation"]),
        ("N° OF", d["n_of"]),
        ("Quantité", d["qte"]),
        ("Annotation série", d["annotation"]),
        ("Statut", d["statut"]),
        ("Confiance", f"{round((d['confidence'] or 0) * 100)} %"),
        ("Date", d["date_creation"].strftime("%d/%m/%Y %H:%M")),
    ]:
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 2, "" if val is None else str(val))
        row += 1

    def section(title: str, headers: list, rows: list) -> None:
        nonlocal row
        row += 1
        ws.cell(row, 1, title).font = Font(bold=True, color=_BLUE)
        row += 1
        for ci, hh in enumerate(headers, 1):
            cell = ws.cell(row, ci, hh)
            cell.font = white_bold
            cell.fill = fill
        row += 1
        for r in rows:
            for ci, val in enumerate(r, 1):
                ws.cell(row, ci, "" if val is None else val)
            row += 1

    section("OPÉRATIONS", OP_HEADERS, [_op_cells(o) for o in d["operations"]])
    section("CONTRÔLES", CTRL_HEADERS, [_ctrl_cells(c) for c in d["controls"]])

    ws.cell(row + 1, 1, "Propulsé par Farness · Agilink Fiches Suiveuses").font = Font(italic=True, size=9, color="9CA3AF")

    for i, w in enumerate([5, 40, 7, 8, 6, 8, 8, 18, 11], 1):
        ws.column_dimensions[chr(64 + i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- PDF --------------------------------------------------------------------


def _lat(s: object) -> str:
    """fpdf2 core fonts are latin-1; map the few chars our form uses outside it."""
    return (
        str("" if s is None else s)
        .replace("œ", "oe")
        .replace("—", "-")
        .replace("→", "->")
        .encode("latin-1", "replace")
        .decode("latin-1")
    )


def to_pdf(fiche: Fiche) -> bytes:
    from fpdf import FPDF

    d = _data(fiche)
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()

    if _LOGO.exists():
        pdf.image(str(_LOGO), x=10, y=9, w=38)
    pdf.set_xy(52, 12)
    pdf.set_font("Helvetica", "B", 15)
    pdf.set_text_color(14, 91, 117)
    pdf.cell(0, 8, _lat("Fiche Suiveuse — Traçabilité Interne"), ln=1)

    pdf.set_text_color(0, 0, 0)
    pdf.set_xy(10, 28)
    pdf.set_font("Helvetica", "", 10)
    info = (
        f"Réf. Produit: {d['ref_produit']}    N° OF: {d['n_of']}    "
        f"Qté: {d['qte']}    Statut: {d['statut']}    "
        f"Confiance: {round((d['confidence'] or 0) * 100)}%    "
        f"Date: {d['date_creation'].strftime('%d/%m/%Y %H:%M')}"
    )
    pdf.cell(0, 6, _lat(info), ln=1)
    pdf.ln(3)

    def table(title: str, headers: list, widths: list, rows: list) -> None:
        pdf.set_font("Helvetica", "B", 11)
        pdf.set_text_color(14, 91, 117)
        pdf.cell(0, 7, _lat(title), ln=1)
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_text_color(255, 255, 255)
        pdf.set_fill_color(14, 91, 117)
        for h, w in zip(headers, widths):
            pdf.cell(w, 6, _lat(h), border=1, align="C", fill=True)
        pdf.ln()
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Helvetica", "", 8)
        for r in rows:
            for val, w in zip(r, widths):
                txt = _lat(val)
                while txt and pdf.get_string_width(txt) > w - 2:
                    txt = txt[:-1]
                pdf.cell(w, 5.5, txt, border=1)
            pdf.ln()
        pdf.ln(3)

    table("OPÉRATIONS", OP_HEADERS, [10, 74, 14, 16, 12, 16, 16, 34, 22],
          [_op_cells(o) for o in d["operations"]])
    table("CONTRÔLES", CTRL_HEADERS, [90, 30, 34, 26], [_ctrl_cells(c) for c in d["controls"]])

    if d["validation"]:
        pdf.set_font("Helvetica", "I", 8)
        pdf.set_text_color(180, 120, 0)
        pdf.multi_cell(0, 5, _lat(f"⚠ {len(d['validation'])} champ(s) à vérifier lors de la validation."))

    # "Powered by Farness" footer — travels with any shared/printed fiche.
    pdf.set_auto_page_break(auto=False)
    if _FARNESS.exists():
        lw = 15
        pdf.image(str(_FARNESS), x=(pdf.w - lw) / 2, y=pdf.h - 15, w=lw)
    pdf.set_y(-9)
    pdf.set_font("Helvetica", "I", 6.5)
    pdf.set_text_color(170, 170, 170)
    pdf.cell(0, 4, _lat("Propulsé par Farness  ·  Agilink — Traçabilité Interne"), align="C")

    out = pdf.output()
    return bytes(out)


# --- CSV --------------------------------------------------------------------


def to_csv(fiche: Fiche) -> bytes:
    d = _data(fiche)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Réf. Produit", d["ref_produit"]])
    w.writerow(["N° OF", d["n_of"]])
    w.writerow(["Quantité", d["qte"]])
    w.writerow(["Statut", d["statut"]])
    w.writerow(["Date", d["date_creation"].strftime("%d/%m/%Y %H:%M")])
    w.writerow([])
    w.writerow(["Partie", *OP_HEADERS])
    for o in d["operations"]:
        w.writerow([o.get("partie"), *_op_cells(o)])
    w.writerow([])
    w.writerow(CTRL_HEADERS)
    for c in d["controls"]:
        w.writerow(_ctrl_cells(c))
    w.writerow([])
    w.writerow(["Propulsé par Farness · Agilink Fiches Suiveuses"])
    # utf-8 BOM so Excel opens the accents correctly
    return buf.getvalue().encode("utf-8-sig")
